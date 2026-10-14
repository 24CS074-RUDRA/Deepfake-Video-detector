from __future__ import annotations

import csv
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "Deepfake_Lens_Project_Presentation_Updated.pptx"
REPORTS = ROOT / "reports"

NAVY = RGBColor(24, 37, 47)
INK = RGBColor(246, 244, 238)
PAPER = RGBColor(255, 253, 249)
TEAL = RGBColor(39, 120, 107)
MINT = RGBColor(60, 141, 112)
CORAL = RGBColor(183, 93, 77)
AMBER = RGBColor(169, 121, 54)
MUTED = RGBColor(99, 111, 105)
LINE = RGBColor(218, 220, 211)


def load_json(path: Path, default=None):
    if not path.exists():
        return default if default is not None else {}
    return json.loads(path.read_text(encoding="utf-8"))


def load_csv(path: Path):
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def add_text(slide, text, x, y, w, h, size=18, color=NAVY, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, color=NAVY, gap=5):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"- {item}"
        paragraph.level = 0
        paragraph.space_after = Pt(gap)
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = color
    return box


def add_box(slide, x, y, w, h, fill=PAPER, line=LINE, radius=False):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    return shape


def add_title(slide, kicker, title, subtitle=None, dark=False):
    color = INK if dark else NAVY
    add_text(slide, kicker.upper(), 0.55, 0.35, 3.8, 0.25, 10, TEAL if not dark else MINT, True, "Aptos Mono")
    add_text(slide, title, 0.55, 0.68, 12.1, 0.62, 28, color, True, "Aptos Display")
    if subtitle:
        add_text(slide, subtitle, 0.58, 1.38, 11.5, 0.38, 12, MUTED if not dark else RGBColor(198, 211, 205))


def add_footer(slide, number, dark=False):
    color = RGBColor(180, 195, 190) if dark else MUTED
    add_text(slide, "DEEPFAKE LENS  /  EXPLAINABLE FACE-SWAP DETECTION", 0.55, 7.12, 7.0, 0.18, 8, color, False, "Aptos Mono")
    add_text(slide, f"{number:02d}", 12.55, 7.08, 0.45, 0.22, 10, color, True, "Aptos Mono", PP_ALIGN.RIGHT)


def add_metric(slide, value, label, x, y, color=TEAL, dark=False):
    add_text(slide, value, x, y, 1.7, 0.42, 25, color, True, "Aptos Display")
    add_text(slide, label.upper(), x, y + 0.45, 2.05, 0.27, 9, RGBColor(192, 204, 198) if dark else MUTED, True, "Aptos Mono")


def add_image(slide, path, x, y, w=None, h=None):
    if path.exists():
        return slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
    return None


def create_deck():
    metrics = load_json(REPORTS / "full_v3" / "test_metrics.json")
    failure_cases = load_json(REPORTS / "full_v3" / "failure_cases.json", [])
    ablations = load_csv(REPORTS / "ablation_val.csv")
    branch_ablation = load_csv(REPORTS / "full_v3" / "test_branch_ablation.csv")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    add_text(slide, "SMART INDIA HACKATHON  /  SGP-II", 0.72, 0.62, 7, 0.25, 11, MINT, True, "Aptos Mono")
    add_text(slide, "Deepfake Lens", 0.68, 1.38, 8.3, 0.9, 48, INK, True, "Aptos Display")
    add_text(slide, "Explainable face-swap deepfake detection", 0.74, 2.42, 7.8, 0.46, 23, RGBColor(218, 229, 223), False, "Aptos Display")
    add_text(slide, "A video-level forensic screening system that combines spatial appearance, frequency-domain artifacts, and temporal consistency into one inspectable verdict.", 0.76, 3.25, 6.0, 1.0, 16, RGBColor(191, 208, 201))
    add_box(slide, 8.25, 1.28, 4.05, 4.58, RGBColor(35, 57, 64), RGBColor(77, 133, 120), True)
    add_text(slide, "THE SIGNAL", 8.68, 1.7, 2.5, 0.25, 10, MINT, True, "Aptos Mono")
    add_text(slide, "3", 8.68, 2.08, 1.2, 0.65, 40, INK, True, "Aptos Display")
    add_text(slide, "evidence branches", 9.78, 2.29, 1.95, 0.25, 13, RGBColor(216, 230, 224))
    add_text(slide, "0.893", 8.68, 3.05, 1.7, 0.5, 28, INK, True, "Aptos Display")
    add_text(slide, "held-out AUC", 10.35, 3.23, 1.45, 0.25, 12, RGBColor(216, 230, 224))
    add_text(slide, "0.802", 8.68, 4.05, 1.7, 0.5, 28, CORAL, True, "Aptos Display")
    add_text(slide, "held-out F1", 10.35, 4.23, 1.45, 0.25, 12, RGBColor(216, 230, 224))
    add_text(slide, "Detection is paired with Grad-CAM, FFT heatmaps, branch attribution, and a suspicion timeline.", 8.68, 5.02, 2.9, 0.58, 12, RGBColor(191, 208, 201))
    add_footer(slide, 1, True)

    # 2. Executive summary
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "01 / Executive summary", "What this project delivers", "A complete research prototype from raw video to explainable web review.")
    cards = [("INPUT", "MP4 / MOV / AVI / WebM", "Upload a video to the local FastAPI service."), ("ANALYSIS", "10 sampled face frames", "Spatial, FFT, and temporal signals are fused."), ("OUTPUT", "REAL or FAKE", "Score, confidence, evidence frames, and forensic JSON."), ("REVIEW", "Human-readable evidence", "The model explains which signals drove suspicion.")]
    for i, (head, value, copy) in enumerate(cards):
        x = 0.65 + i * 3.08
        add_box(slide, x, 2.15, 2.72, 2.05, RGBColor(250, 249, 244), LINE, True)
        add_text(slide, head, x + 0.22, 2.4, 2.2, 0.2, 9, TEAL, True, "Aptos Mono")
        add_text(slide, value, x + 0.22, 2.78, 2.25, 0.52, 18, NAVY, True, "Aptos Display")
        add_text(slide, copy, x + 0.22, 3.48, 2.2, 0.47, 11, MUTED)
    add_bullets(slide, ["Built with PyTorch, TorchVision, OpenCV, RetinaFace, FastAPI, scikit-learn, and optional SHAP.", "Designed for forensic assistance: a model score is evidence for review, not proof of authenticity.", "The held-out test report contains 300 videos: 150 real and 150 fake, with a 50 / 50 confusion-matrix support."] , 0.85, 4.85, 11.7, 1.25, 15)
    add_footer(slide, 2)

    # 3. Problem and objectives
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(239, 244, 239)
    add_title(slide, "02 / Motivation", "Why face-swap detection needs more than one clue", "Spatial artifacts can be subtle, frequency artifacts can be hidden, and temporal glitches may appear only across frames.")
    add_box(slide, 0.7, 2.12, 5.85, 3.85, NAVY, NAVY, True)
    add_text(slide, "THE CHALLENGE", 1.02, 2.46, 2.5, 0.22, 10, MINT, True, "Aptos Mono")
    add_text(slide, "A convincing fake can look real in one frame.", 1.02, 2.92, 4.7, 0.7, 27, INK, True, "Aptos Display")
    add_text(slide, "Compression, lighting, pose, skin texture, and identity-transfer artifacts interact. A robust detector should look at complementary representations and preserve enough evidence for a reviewer to question the result.", 1.04, 4.06, 4.8, 1.05, 15, RGBColor(205, 220, 214))
    add_text(slide, "OBJECTIVES", 7.05, 2.25, 2.5, 0.22, 10, TEAL, True, "Aptos Mono")
    add_bullets(slide, ["Detect face-swap videos at video level.", "Combine RGB appearance, FFT texture, and sequence dynamics.", "Evaluate with held-out test metrics and controlled ablations.", "Generate explanations that support human review.", "Expose the complete workflow through a local upload application."], 7.05, 2.68, 5.35, 2.85, 18)
    add_text(slide, "Success is therefore both predictive and inspectable.", 7.05, 5.7, 5.2, 0.35, 17, CORAL, True, "Aptos Display")
    add_footer(slide, 3)

    # 4. Dataset and data split
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "03 / Data", "From FaceForensics++ videos to sequence samples", "Raw videos are separated into real and fake folders; generated artifacts are excluded from version control.")
    add_box(slide, 0.7, 2.05, 3.4, 3.7, RGBColor(246, 248, 242), LINE, True)
    add_text(slide, "RAW DATA", 1.0, 2.36, 2.0, 0.22, 10, TEAL, True, "Aptos Mono")
    add_text(slide, "ffpp_real/*.mp4", 1.0, 2.84, 2.5, 0.33, 18, NAVY, True, "Aptos Mono")
    add_text(slide, "ffpp_fake/*.mp4", 1.0, 3.32, 2.5, 0.33, 18, CORAL, True, "Aptos Mono")
    add_text(slide, "Video IDs are grouped before frame rows are assembled, preventing frames from a single video from being treated as independent examples.", 1.0, 4.15, 2.65, 1.05, 13, MUTED)
    add_box(slide, 4.65, 2.05, 3.4, 3.7, RGBColor(246, 248, 242), LINE, True)
    add_text(slide, "HELD-OUT TEST", 4.95, 2.36, 2.4, 0.22, 10, TEAL, True, "Aptos Mono")
    add_text(slide, "300", 4.95, 2.82, 1.45, 0.5, 30, NAVY, True, "Aptos Display")
    add_text(slide, "videos evaluated", 6.1, 3.02, 1.5, 0.22, 11, MUTED)
    add_text(slide, "150 real", 4.95, 3.72, 1.45, 0.32, 17, MINT, True, "Aptos Display")
    add_text(slide, "150 fake", 6.55, 3.72, 1.2, 0.32, 17, CORAL, True, "Aptos Display")
    add_text(slide, "Full dataset: 2,000 videos = 1,000 real + 1,000 fake. Split: 1,400 train, 300 validation, 300 held-out test.", 4.95, 4.48, 2.65, 0.72, 12, MUTED)
    add_box(slide, 8.6, 2.05, 4.0, 3.7, NAVY, NAVY, True)
    add_text(slide, "PER VIDEO", 8.95, 2.36, 2.0, 0.22, 10, MINT, True, "Aptos Mono")
    add_text(slide, "10", 8.95, 2.83, 1.25, 0.5, 30, INK, True, "Aptos Display")
    add_text(slide, "sampled sequence frames", 10.05, 3.03, 2.05, 0.22, 12, RGBColor(210, 225, 219))
    add_text(slide, "RGB face crop + grayscale FFT image + label", 8.95, 3.85, 2.85, 0.55, 17, INK, True, "Aptos Display")
    add_text(slide, "Metadata fields include image_path, fft_path, label, video_id, frame_idx, and split.", 8.95, 4.7, 2.9, 0.58, 12, RGBColor(196, 211, 205))
    add_footer(slide, 4)

    # 5. Pipeline
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(239, 244, 239)
    add_title(slide, "04 / Preprocessing", "The evidence pipeline", "Each raw video is converted into aligned spatial and frequency inputs before model training.")
    steps = [("01", "Extract frames", "OpenCV\nFRAME_INTERVAL = 30"), ("02", "Detect faces", "RetinaFace\nlargest face crop"), ("03", "Transform", "Resize to 224 x 224\nRGB normalization"), ("04", "FFT", "Log magnitude\nfrequency image"), ("05", "Metadata", "Video-level split\nsequence rows")]
    for i, (num, head, copy) in enumerate(steps):
        x = 0.7 + i * 2.48
        add_box(slide, x, 2.35, 1.98, 2.35, PAPER, LINE, True)
        add_text(slide, num, x + 0.2, 2.62, 0.42, 0.25, 12, TEAL, True, "Aptos Mono")
        add_text(slide, head, x + 0.2, 3.08, 1.55, 0.42, 17, NAVY, True, "Aptos Display")
        add_text(slide, copy, x + 0.2, 3.75, 1.52, 0.65, 11, MUTED, False, "Aptos Mono")
        if i < 4:
            add_text(slide, ">", x + 2.1, 3.2, 0.3, 0.3, 20, CORAL, True, "Aptos Display", PP_ALIGN.CENTER)
    add_text(slide, "Failure handling", 0.8, 5.45, 1.5, 0.23, 11, CORAL, True, "Aptos Mono")
    add_text(slide, "Unreadable image rows are replaced with blank tensors so a single corrupt artifact does not crash sequence construction. Training applies horizontal flip, color jitter, and occasional JPEG compression augmentation.", 2.25, 5.39, 9.7, 0.45, 13, MUTED)
    add_footer(slide, 5)

    # 6. Architecture
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    add_title(slide, "05 / Model", "Three complementary branches, one fused verdict", "The architecture returns the fused logit plus branch-level and per-frame signals for inspection.", True)
    branches = [("RGB / FRAME", "EfficientNet-B0", "1280-d features", TEAL), ("FFT / FREQUENCY", "3-layer CNN", "128-d features", AMBER), ("SEQUENCE / TEMPORAL", "Bidirectional LSTM", "256-d embedding", CORAL)]
    for i, (head, model, dim, color) in enumerate(branches):
        x = 0.75 + i * 4.02
        add_box(slide, x, 2.12, 3.35, 1.75, RGBColor(35, 57, 64), color, True)
        add_text(slide, head, x + 0.25, 2.38, 2.8, 0.22, 10, color, True, "Aptos Mono")
        add_text(slide, model, x + 0.25, 2.78, 2.8, 0.35, 18, INK, True, "Aptos Display")
        add_text(slide, dim, x + 0.25, 3.32, 2.8, 0.22, 12, RGBColor(201, 218, 211))
        if i < 2: add_text(slide, "+", x + 3.5, 2.72, 0.35, 0.4, 24, MINT, True, "Aptos Display", PP_ALIGN.CENTER)
    add_text(slide, "FUSION", 4.96, 4.25, 2.0, 0.22, 10, MINT, True, "Aptos Mono", PP_ALIGN.CENTER)
    add_box(slide, 4.05, 4.58, 5.2, 1.02, RGBColor(53, 78, 80), MINT, True)
    add_text(slide, "mean pooled branch features  ->  Linear(256)  ->  ReLU  ->  Dropout(0.3)  ->  Linear(1)", 4.35, 4.89, 4.6, 0.3, 13, INK, True, "Aptos Mono", PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "Auxiliary heads", 0.85, 6.08, 1.4, 0.2, 11, RGBColor(201, 218, 211), True, "Aptos Mono")
    add_text(slide, "Frame and frequency branches also produce per-frame logits; the temporal branch produces sequence-level logits. The output exposes embeddings for attribution.", 2.25, 6.02, 9.8, 0.36, 13, RGBColor(201, 218, 211))
    add_footer(slide, 6, True)

    # 7. Training
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "06 / Training", "Training strategy and reproducible outputs", "The scripts package training, checkpointing, curve generation, evaluation, ablation, and explanation as runnable commands.")
    add_box(slide, 0.7, 2.05, 5.45, 3.95, RGBColor(247, 248, 243), LINE, True)
    add_text(slide, "DEFAULT FULL RUN", 1.0, 2.36, 2.5, 0.2, 10, TEAL, True, "Aptos Mono")
    add_text(slide, "10 epochs  /  batch 4  /  lr 1e-4  /  sequence length 10", 1.0, 2.78, 4.3, 0.35, 17, NAVY, True, "Aptos Display")
    add_bullets(slide, ["BCEWithLogitsLoss with positive-class weighting.", "AdamW optimizer with a lower learning rate for the EfficientNet backbone.", "Backbone frozen initially; unfrozen at epoch 3.", "ReduceLROnPlateau scheduler and early stopping after 5 stale epochs.", "Best checkpoint selected by validation AUC."], 1.0, 3.45, 4.65, 2.0, 14)
    add_box(slide, 6.65, 2.05, 5.95, 3.95, RGBColor(247, 248, 243), LINE, True)
    add_text(slide, "ARTIFACTS", 6.95, 2.36, 2.5, 0.2, 10, TEAL, True, "Aptos Mono")
    artifact_items = ["trained_models/<run>/best.pt", "trained_models/<run>/last.pt", "trained_models/<run>/config.json", "reports/<run>/history.csv", "reports/<run>/test_metrics.json", "ROC, accuracy, loss, confusion matrix, failure report"]
    add_bullets(slide, artifact_items, 6.95, 2.8, 5.0, 2.5, 15)
    add_text(slide, "CPU smoke training is also supported with --limit 5 and two epochs.", 0.85, 6.35, 6.0, 0.25, 12, CORAL, True, "Aptos Mono")
    add_footer(slide, 7)

    # 8. Ablations
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(239, 244, 239)
    add_title(slide, "07 / Ablation study", "What each representation contributes", "Validation ablations compare the same task with frame, frequency, temporal, and full configurations.")
    add_text(slide, "VALIDATION AUC", 0.85, 2.08, 2.0, 0.22, 10, TEAL, True, "Aptos Mono")
    max_auc = max([float(row["val_auc"]) for row in ablations] or [1.0])
    for i, row in enumerate(ablations):
        name = row["configuration"]
        auc = float(row["val_auc"])
        y = 2.65 + i * 0.72
        add_text(slide, name, 0.85, y, 1.65, 0.25, 14, NAVY, True, "Aptos Mono")
        add_box(slide, 2.65, y + 0.02, 6.0, 0.27, RGBColor(223, 231, 223), RGBColor(223, 231, 223), True)
        add_box(slide, 2.65, y + 0.02, 6.0 * auc / max_auc, 0.27, TEAL if name == "full" else RGBColor(111, 160, 137), RGBColor(111, 160, 137), True)
        add_text(slide, f"{auc:.3f}", 8.9, y - 0.03, 0.75, 0.3, 15, NAVY, True, "Aptos Mono", PP_ALIGN.RIGHT)
    add_box(slide, 9.6, 2.15, 2.85, 3.55, NAVY, NAVY, True)
    add_text(slide, "READOUT", 9.95, 2.48, 1.5, 0.22, 10, MINT, True, "Aptos Mono")
    add_text(slide, "Full fusion reaches the strongest validation AUC: 0.945.", 9.95, 2.94, 2.05, 0.75, 21, INK, True, "Aptos Display")
    add_text(slide, "Frame and frame+temporal are also strong. Frequency-only is weaker alone, but adds a distinct signal to the fusion system.", 9.95, 4.1, 2.05, 1.0, 13, RGBColor(201, 218, 211))
    add_text(slide, "Validation values are not the held-out test result.", 0.85, 6.18, 5.2, 0.25, 12, CORAL, True, "Aptos Mono")
    add_footer(slide, 8)

    # 9. Test results
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "08 / Held-out evaluation", "Final test performance", "Metrics from reports/full_v3/test_metrics.json using the operating threshold selected during evaluation.")
    test_metrics = [("0.893", "AUC", TEAL), ("0.780", "Accuracy", NAVY), ("0.802", "F1 score", CORAL), ("0.893", "Recall", MINT)]
    for i, (value, label, color) in enumerate(test_metrics):
        add_box(slide, 0.75 + i * 3.05, 2.05, 2.65, 1.32, RGBColor(248, 249, 245), LINE, True)
        add_metric(slide, value, label, 1.0 + i * 3.05, 2.35, color)
    add_image(slide, REPORTS / "full_v3" / "roc_curve.png", 0.8, 3.72, w=4.0)
    add_image(slide, REPORTS / "full_v3" / "confusion_matrix.png", 4.95, 3.72, w=3.4)
    add_box(slide, 8.65, 3.7, 3.75, 2.52, RGBColor(247, 248, 243), LINE, True)
    add_text(slide, "CONFUSION MATRIX", 8.98, 4.02, 2.1, 0.2, 10, TEAL, True, "Aptos Mono")
    add_text(slide, "100  true real", 8.98, 4.48, 2.4, 0.3, 18, MINT, True, "Aptos Mono")
    add_text(slide, "134  true fake", 8.98, 4.92, 2.4, 0.3, 18, CORAL, True, "Aptos Mono")
    add_text(slide, "50 false positives  /  16 false negatives\nEER: 0.193  /  Precision: 0.728", 8.98, 5.4, 2.8, 0.55, 12, MUTED, False, "Aptos Mono")
    add_footer(slide, 9)

    # 10. Branch test ablation
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(239, 244, 239)
    add_title(slide, "09 / Test branch analysis", "Fusion is the best operating configuration", "Held-out branch comparison from full_v3 shows how the fused score behaves against individual signals.")
    headers = ["BRANCH", "ACCURACY", "F1", "AUC"]
    values = [(row["branch"], row["accuracy"], row["f1"], row["auc"]) for row in branch_ablation]
    x0, y0 = 0.9, 2.25
    widths = [3.1, 2.1, 2.1, 2.1]
    for j, header in enumerate(headers):
        add_box(slide, x0 + sum(widths[:j]), y0, widths[j], 0.58, NAVY, NAVY)
        add_text(slide, header, x0 + sum(widths[:j]) + 0.16, y0 + 0.18, widths[j] - 0.3, 0.2, 10, INK, True, "Aptos Mono")
    for i, row in enumerate(values):
        y = y0 + 0.58 + i * 0.68
        fill = RGBColor(247, 248, 243) if i % 2 == 0 else PAPER
        for j, value in enumerate(row):
            add_box(slide, x0 + sum(widths[:j]), y, widths[j], 0.68, fill, LINE)
            text = value if j == 0 else f"{float(value):.3f}"
            add_text(slide, text, x0 + sum(widths[:j]) + 0.16, y + 0.2, widths[j] - 0.3, 0.22, 15 if j == 0 else 14, CORAL if row[0] == "fusion" and j > 0 else NAVY, row[0] == "fusion", "Aptos Mono")
    add_box(slide, 9.75, 2.25, 2.5, 3.3, NAVY, NAVY, True)
    add_text(slide, "TAKEAWAY", 10.05, 2.6, 1.5, 0.2, 10, MINT, True, "Aptos Mono")
    add_text(slide, "Fusion improves accuracy and F1 over every single branch.", 10.05, 3.08, 1.8, 0.9, 22, INK, True, "Aptos Display")
    add_text(slide, "Frequency is complementary rather than dominant: its standalone AUC is 0.672, while the fused model reaches 0.893.", 10.05, 4.35, 1.8, 0.76, 12, RGBColor(201, 218, 211))
    add_footer(slide, 10)

    # 11. Explainability
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "10 / Explainability", "From a score to inspectable evidence", "The explanation pipeline ranks suspicious frames and attributes the fused signal to model branches.")
    image_paths = [REPORTS / "explanations" / "000" / "frame_0000_gradcam.jpg", REPORTS / "explanations" / "000" / "frame_0004_gradcam.jpg", REPORTS / "explanations" / "000" / "frame_0008_gradcam.jpg"]
    for i, path in enumerate(image_paths):
        add_image(slide, path, 0.8 + i * 2.08, 2.1, w=1.85, h=1.5)
        add_text(slide, f"Grad-CAM frame {i + 1}", 0.8 + i * 2.08, 3.75, 1.85, 0.2, 10, MUTED, False, "Aptos Mono", PP_ALIGN.CENTER)
    add_box(slide, 7.35, 2.1, 5.2, 3.6, RGBColor(247, 248, 243), LINE, True)
    add_text(slide, "EXPLANATION OUTPUTS", 7.7, 2.43, 2.8, 0.2, 10, TEAL, True, "Aptos Mono")
    add_bullets(slide, ["Grad-CAM overlays: top-n high-suspicion sampled frames.", "FFT heatmaps: frequency-domain artifact visualization.", "Branch contribution chart: SHAP KernelExplainer when available; fusion-ablation fallback otherwise.", "Suspicion timeline: per-frame sigmoid scores for each branch.", "explanation.json: device, score, branch scores, overlays, contributions, and timeline."], 7.7, 2.85, 4.35, 2.35, 14)
    add_text(slide, "Interpretation remains assistive evidence, not standalone proof.", 0.85, 5.82, 6.1, 0.3, 17, CORAL, True, "Aptos Display")
    add_footer(slide, 11)

    # 12. Web application
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(239, 244, 239)
    add_title(slide, "11 / Product surface", "Deepfake Lens: local forensic review workspace", "A responsive single-page HTML frontend is served by FastAPI and talks to the active full_v3 checkpoint.")
    add_box(slide, 0.75, 2.02, 5.0, 3.95, NAVY, NAVY, True)
    add_text(slide, "UPLOAD", 1.08, 2.38, 1.3, 0.2, 10, MINT, True, "Aptos Mono")
    add_text(slide, "Drop a video here", 1.08, 2.84, 3.8, 0.4, 24, INK, True, "Aptos Display")
    add_text(slide, "MP4, MOV, AVI, or WebM", 1.08, 3.38, 3.4, 0.25, 13, RGBColor(201, 218, 211))
    add_box(slide, 1.08, 4.1, 3.1, 0.52, TEAL, TEAL, True)
    add_text(slide, "Analyze video", 1.08, 4.23, 3.1, 0.2, 13, INK, True, "Aptos Mono", PP_ALIGN.CENTER)
    add_text(slide, "Local processing keeps the project self-contained during review.", 1.08, 5.12, 3.85, 0.4, 12, RGBColor(201, 218, 211))
    add_text(slide, "API FLOW", 6.35, 2.16, 1.5, 0.2, 10, TEAL, True, "Aptos Mono")
    api_steps = [("POST /api/analyze", "Validate type -> save temporary upload"), ("scripts.explain.run", "Faces -> RGB/FFT tensors -> model"), ("Forensic report", "Verdict, confidence, findings, evidence"), ("GET /assets/...", "Serve Grad-CAM and JSON artifacts")]
    for i, (head, copy) in enumerate(api_steps):
        y = 2.62 + i * 0.78
        add_box(slide, 6.35, y, 5.65, 0.58, PAPER, LINE, True)
        add_text(slide, head, 6.6, y + 0.12, 2.3, 0.2, 12, NAVY, True, "Aptos Mono")
        add_text(slide, copy, 8.95, y + 0.12, 2.75, 0.2, 11, MUTED)
    add_text(slide, "UI result sections: model verdict, confidence meter, AI analysis, forensic report, detected frame evidence, branch contribution, and suspicion timeline.", 6.38, 5.95, 5.45, 0.43, 13, CORAL, True, "Aptos Display")
    add_footer(slide, 12)

    # 13. Setup and usage
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "12 / Reproducibility", "How to run the project", "Windows PowerShell workflow documented in README.md.")
    commands = [
        ("1. Environment", "py -3.13 -m venv venv\nvenv\\Scripts\\activate\npip install -r requirements.txt"),
        ("2. Preprocess", "python -m scripts.preprocess --limit 5\npython -m scripts.inspect_dataset"),
        ("3. Train", "python -m scripts.train --epochs 10 --batch-size 4 --lr 1e-4 --branches full --seq-len 10 --run-name full_model"),
        ("4. Evaluate", "python -m scripts.evaluate --run-name full_model\npython -m pytest -q"),
        ("5. Serve", "venv\\Scripts\\python.exe -m uvicorn ai.api:app --host 127.0.0.1 --port 8000"),
    ]
    for i, (head, command) in enumerate(commands):
        x = 0.72 + (i % 3) * 4.12
        y = 2.05 + (i // 3) * 2.05
        add_box(slide, x, y, 3.55, 1.58, RGBColor(247, 248, 243), LINE, True)
        add_text(slide, head, x + 0.24, y + 0.22, 2.85, 0.22, 11, TEAL, True, "Aptos Mono")
        add_text(slide, command, x + 0.24, y + 0.62, 3.05, 0.7, 11, NAVY, False, "Aptos Mono")
    add_text(slide, "Open http://127.0.0.1:8000 and upload a supported video.", 0.85, 6.45, 6.5, 0.25, 14, CORAL, True, "Aptos Mono")
    add_footer(slide, 13)

    # 14. Limitations
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor(239, 244, 239)
    add_title(slide, "13 / Responsible use", "Known limitations and risks", "The project is a strong research prototype, but its operating context matters.")
    limitations = [
        ("DATA SCALE", "Metrics can be unreliable when only a few test videos are available; results should be rechecked on broader, diverse data."),
        ("CPU BOTTLENECK", "RetinaFace face extraction is CPU-bound on native Windows TensorFlow installations and can be slow."),
        ("ATTRIBUTION", "SHAP is optional; when unavailable, the system uses a fusion-ablation attribution fallback."),
        ("THRESHOLD", "The operating threshold is selected from evaluation data and confidence is not a guarantee of correctness."),
        ("DOMAIN SHIFT", "Compression, identities, manipulations, lighting, and camera conditions outside the data distribution may reduce accuracy."),
        ("HUMAN REVIEW", "Outputs are assistive screening evidence, not a standalone authenticity or identity claim."),
    ]
    for i, (head, copy) in enumerate(limitations):
        x = 0.78 + (i % 3) * 4.1
        y = 2.08 + (i // 3) * 1.85
        add_box(slide, x, y, 3.55, 1.42, PAPER, LINE, True)
        add_text(slide, head, x + 0.22, y + 0.24, 2.9, 0.2, 10, CORAL, True, "Aptos Mono")
        add_text(slide, copy, x + 0.22, y + 0.58, 3.0, 0.62, 12, NAVY)
    add_footer(slide, 14)

    # 15. Future work and conclusion
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
    add_title(slide, "14 / Next steps", "A foundation for stronger forensic tooling", "The current system creates a credible baseline; the next improvements should target generalization, speed, and evaluation rigor.", True)
    add_bullets(slide, ["Expand evaluation across manipulation methods, identities, compression levels, and unseen sources.", "Add calibration, confidence intervals, and cross-dataset testing so the score communicates uncertainty.", "Accelerate face extraction with WSL2, TensorFlow-DirectML, batching, or a faster detector.", "Add robust tracking and multi-face handling for realistic videos.", "Turn reports into an auditable review workflow with versioned models and evidence bundles."], 0.9, 2.18, 6.45, 3.45, 18, INK, 8)
    add_box(slide, 8.3, 2.1, 3.9, 3.75, RGBColor(35, 57, 64), RGBColor(77, 133, 120), True)
    add_text(slide, "THE PROJECT IN ONE LINE", 8.67, 2.48, 2.7, 0.22, 10, MINT, True, "Aptos Mono")
    add_text(slide, "Detect the signal.\nShow the evidence.\nKeep a human in the loop.", 8.67, 3.05, 2.9, 1.45, 25, INK, True, "Aptos Display")
    add_text(slide, "Deepfake Lens combines a three-branch model, held-out evaluation, and a usable forensic interface into one end-to-end project.", 8.67, 5.0, 2.9, 0.55, 13, RGBColor(201, 218, 211))
    add_footer(slide, 15, True)

    # 16. Appendix / traceability
    slide = prs.slides.add_slide(blank); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = PAPER
    add_title(slide, "Appendix", "Project map and evidence sources", "Every major claim in this deck is grounded in an implementation module or stored report artifact.")
    rows = [
        ("Model", "ai/models/deepfake_detector.py", "EfficientNet-B0, CNN, BiLSTM, fusion"),
        ("Dataset", "ai/models/dataset.py", "Video grouping, sampling, transforms"),
        ("Preprocessing", "ai/preprocessing/", "Frames, faces, FFT, metadata"),
        ("Training", "scripts/train.py", "Loss, optimizer, scheduler, checkpoints"),
        ("Evaluation", "scripts/evaluate.py", "Metrics, ROC, confusion matrix, failures"),
        ("Explainability", "ai/explain/ and scripts/explain.py", "Grad-CAM, FFT heatmaps, SHAP/fallback"),
        ("Application", "ai/api.py and frontend/index.html", "Upload API and review console"),
        ("Validation", "tests/", "Aggregation, pipeline, model, Grad-CAM, SHAP tests"),
    ]
    widths = [1.6, 4.35, 5.95]
    x0, y0 = 0.75, 2.0
    for j, header in enumerate(["AREA", "SOURCE", "EVIDENCE"]):
        add_box(slide, x0 + sum(widths[:j]), y0, widths[j], 0.48, NAVY, NAVY)
        add_text(slide, header, x0 + sum(widths[:j]) + 0.14, y0 + 0.14, widths[j] - 0.2, 0.2, 9, INK, True, "Aptos Mono")
    for i, row in enumerate(rows):
        y = y0 + 0.48 + i * 0.49
        for j, value in enumerate(row):
            add_box(slide, x0 + sum(widths[:j]), y, widths[j], 0.49, RGBColor(248, 249, 245) if i % 2 == 0 else PAPER, LINE)
            add_text(slide, value, x0 + sum(widths[:j]) + 0.14, y + 0.14, widths[j] - 0.22, 0.18, 10 if j else 11, NAVY if j else TEAL, j == 0, "Aptos Mono" if j == 1 else "Aptos")
    add_text(slide, "Generated from repository state on 2026-10-14.", 0.85, 6.55, 4.5, 0.22, 10, MUTED, False, "Aptos Mono")
    add_footer(slide, 16)

    prs.core_properties.title = "Deepfake Lens - Explainable Face-Swap Deepfake Detection"
    prs.core_properties.subject = "Complete project presentation generated from repository artifacts"
    prs.core_properties.author = "Deepfake Lens project"
    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    create_deck()